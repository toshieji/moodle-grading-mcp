#!/usr/bin/env python3
"""cloud/grading_job.py の提出ファイル読み取りと「読めない提出は採点しない」の検証。

依存なしで動く（python tests/test_grading_job.py）。Moodle と Anthropic には接続しない。
ダウンロード・採点・保存を差し替え、pptx などを実際に生成して通す。失敗すると終了コード 1。
"""
from __future__ import annotations

import io
import os
import sys
import types

os.environ.setdefault("MOODLE_URL", "https://moodle.example.invalid")
os.environ.setdefault("MOODLE_TOKEN", "test-token")
os.environ.setdefault("ANTHROPIC_API_KEY", "test-key")
os.environ["MOODLE_ALLOW_WRITE"] = "1"
os.environ["MOODLE_WRITE_COURSE_ALLOWLIST"] = "900"
os.environ["GRADE_COURSE_IDS"] = "900"

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "cloud"))
import grading_job as gj  # noqa: E402

FAILED: list[str] = []


def check(cond: bool, label: str) -> None:
    print(("  ok   " if cond else "  FAIL ") + label)
    if not cond:
        FAILED.append(label)


def png_bytes() -> bytes:
    from PIL import Image as PILImage

    buf = io.BytesIO()
    PILImage.new("RGB", (40, 30), (200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def pptx_bytes(text: str | None, with_image: bool) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if text:
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = text
    if with_image:
        slide.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(3))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def run_extract(blobs: dict[str, bytes | Exception]):
    """filename -> 中身（例外ならダウンロード失敗）で extract_submitted_files を呼ぶ。"""
    def fake_download(url: str) -> bytes:
        v = blobs[url]
        if isinstance(v, Exception):
            raise v
        return v

    orig = gj._download_file
    gj._download_file = fake_download
    try:
        files = [{"filename": n, "fileurl": n, "filesize": 10} for n in blobs]
        return gj.extract_submitted_files(files)
    finally:
        gj._download_file = orig


def test_extract() -> None:
    print("extract_submitted_files")
    text, unreadable, images = run_extract({"企画レポート.pptx": pptx_bytes("主KPI：CTAクリック数", True)})
    check("主KPI：CTAクリック数" in text, "pptx のスライド本文を読める（従来は未対応で読めず）")
    check(not unreadable, "pptx は読めないファイル扱いにならない")
    check(len(images) == 1 and images[0]["label"].startswith("企画レポート.pptx"),
          "pptx の埋め込み画像をファイル名つきで返す")

    text, unreadable, images = run_extract({"図だけ.pptx": pptx_bytes(None, True)})
    check(not unreadable and len(images) == 1, "文字のない図だけの pptx も画像として採点に回す")

    _, unreadable, images = run_extract({"shot.png": png_bytes()})
    check(not unreadable and len(images) == 1, "画像ファイルそのものを読める")

    _, unreadable, _ = run_extract({"資料.key": b"\x00\x01"})
    check(len(unreadable) == 1 and "未対応" in unreadable[0], "未対応形式は読めないファイルとして報告する")

    _, unreadable, _ = run_extract({"原稿.docx": b"not a zip"})
    check(len(unreadable) == 1, "壊れた docx は読めないファイルとして報告する")

    _, unreadable, _ = run_extract({"原稿.docx": RuntimeError("403")})
    check(len(unreadable) == 1 and "ダウンロード失敗" in unreadable[0], "ダウンロード失敗を報告する")

    orig = gj.MAX_IMAGES_PER_SUBMISSION
    gj.MAX_IMAGES_PER_SUBMISSION = 1
    try:
        _, _, images = run_extract({"a.png": png_bytes(), "b.pptx": pptx_bytes("本文", True)})
        check(len(images) == 1, "画像の上限を提出全体で守る")
    finally:
        gj.MAX_IMAGES_PER_SUBMISSION = orig


def test_user_content() -> None:
    print("build_user_content")
    check(gj.build_user_content("本文", []) == "本文", "画像なしなら文字列のまま")
    blocks = gj.build_user_content("本文", [{"label": "x.pptx / slide1_img1", "data": b"\x89PNG", "format": "png"}])
    kinds = [b["type"] for b in blocks]
    check(kinds == ["text", "text", "image"], "本文・画像ラベル・画像の順に並べる")
    check(blocks[2]["source"]["media_type"] == "image/png", "media_type を形式から付ける")


def test_main_skips_unreadable() -> None:
    """読めない提出は採点も保存もしない（0点を付けない）。読める提出は従来どおり保存する。"""
    print("main: 読めない提出の扱い")
    saved: list[int] = []
    graded_calls: list[int] = []
    subs = {
        1: {"onlinetext": "log", "files": ["a.key"], "file_text": "", "unreadable": ["a.key（未対応の形式 .key）"],
            "images": []},
        2: {"onlinetext": "log", "files": ["b.pptx"], "file_text": "本文", "unreadable": [], "images": []},
    }
    patches = {
        "fetch_rubric_sheet_text": lambda: "rubric",
        "load_ai_usage_rubric": lambda: "ai rubric",
        "list_assignments": lambda cid: [{"id": 10, "cmid": 20, "name": "事前課題3", "grademax": 30}],
        "list_pending": lambda aid: [{"userid": 1}, {"userid": 2}],
        "get_guide_criteria": lambda cmid: None,
        "get_submission": lambda aid, uid: subs[uid],
        "grade_submission": lambda *a, **k: (graded_calls.append(1) or
                                             {"grade": 20, "feedback_html": "<p>x</p>",
                                              "confidence": "high", "needs_human_review": False}),
        "save_grade_draft": lambda aid, uid, *a, **k: saved.append(uid),
    }
    originals = {k: getattr(gj, k) for k in patches}
    sys.modules.setdefault("anthropic", types.SimpleNamespace(Anthropic=lambda api_key: object()))
    for k, v in patches.items():
        setattr(gj, k, v)
    try:
        gj.main()
    finally:
        for k, v in originals.items():
            setattr(gj, k, v)
    check(saved == [2], "読めない提出（user 1）は保存されず、読める提出（user 2）だけ保存される")
    check(len(graded_calls) == 1, "読めない提出は採点モデルを呼ばない")

    # 再採点の指定: 採点済み（pending に出ない）user 3 を指定すると採点し直す。指定外の課題は触らない
    subs[3] = subs[2]
    saved.clear()
    patches["list_pending"] = lambda aid: []
    orig_targets = gj.REGRADE_TARGETS
    gj.REGRADE_TARGETS = {(10, 3), (99, 4)}
    for k, v in patches.items():
        setattr(gj, k, v)
    try:
        gj.main()
    finally:
        gj.REGRADE_TARGETS = orig_targets
        for k, v in originals.items():
            setattr(gj, k, v)
    check(saved == [3], "REGRADE_TARGETS で指定した採点済みの提出だけを採点し直す")


if __name__ == "__main__":
    test_extract()
    test_user_content()
    test_main_skips_unreadable()
    print(f"\n{'FAILED: ' + str(len(FAILED)) if FAILED else 'all passed'}")
    sys.exit(1 if FAILED else 0)
